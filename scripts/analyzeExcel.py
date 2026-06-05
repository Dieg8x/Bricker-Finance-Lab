import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXCEL = Path("/Users/dieg8x/Downloads/Formulario Bricker - Examen Completo Simple.xlsx")
OUT = ROOT / "src" / "data" / "excelExtract.json"

INPUT_FILLS = {"FFF2CC", "00FFF2CC", "FFFFF2CC"}
RESULT_FILLS = {"E2F0D9", "00E2F0D9", "FFE2F0D9", "D9EAF7", "00D9EAF7", "FFD9EAF7"}


def cell_fill(cell):
    fill = cell.fill
    color = fill.fgColor.rgb or fill.fgColor.indexed
    return str(color) if color else None


def json_value(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def left_label(ws, cell):
    candidates = []
    for col in range(max(1, cell.column - 4), cell.column):
        value = ws.cell(cell.row, col).value
        if value not in (None, ""):
            candidates.append(str(value))
    above = ws.cell(max(1, cell.row - 1), cell.column).value
    if above not in (None, ""):
        candidates.append(str(above))
    return candidates[-1] if candidates else ""


def extract():
    wb = load_workbook(EXCEL, data_only=False)
    workbook = {"source": str(EXCEL), "sheetCount": len(wb.sheetnames), "sheets": []}
    for ws in wb.worksheets:
        formulas = []
        possible_inputs = []
        possible_outputs = []
        labels = []
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value in (None, ""):
                    continue
                fill = cell_fill(cell)
                label = left_label(ws, cell)
                normalized = json_value(value)
                if isinstance(normalized, str) and normalized.startswith("="):
                    formulas.append({
                        "cell": cell.coordinate,
                        "formula": normalized,
                        "nearLabel": label,
                    })
                    if fill in RESULT_FILLS:
                        possible_outputs.append({
                            "cell": cell.coordinate,
                            "label": label,
                            "formula": normalized,
                        })
                elif fill in INPUT_FILLS:
                    possible_inputs.append({
                        "cell": cell.coordinate,
                        "label": label or str(normalized),
                        "value": normalized,
                    })
                elif isinstance(normalized, str):
                    labels.append({"cell": cell.coordinate, "text": normalized[:140]})

        workbook["sheets"].append({
            "name": ws.title,
            "dimensions": {"rows": ws.max_row, "columns": ws.max_column},
            "formulaCount": len(formulas),
            "formulas": formulas[:160],
            "possibleInputs": possible_inputs[:80],
            "possibleOutputs": possible_outputs[:80],
            "labels": labels[:80],
        })
    return workbook


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    data = extract()
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Workbook: {data['source']}")
    print(f"Sheets: {data['sheetCount']}")
    for sheet in data["sheets"]:
        print(f"- {sheet['name']}: {sheet['formulaCount']} formulas, {len(sheet['possibleInputs'])} possible inputs")
    print(f"Saved: {OUT}")
